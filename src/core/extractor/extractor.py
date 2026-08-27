"""Metadata Extractor module for TraceLens.

Provides robust, object-oriented metadata extraction across a wide variety
of file formats including PDF, images, audio, office documents, archives,
structured data, source code, databases, videos, and generic binary files.
"""

from __future__ import annotations

import configparser
import csv
import hashlib
import json
import mimetypes
import os
import re
import sqlite3
import tarfile
import wave
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from typing import Any, Callable, Iterable

# Optional third-party dependencies with graceful fallbacks
try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    from PIL import Image, ExifTags
except ImportError:
    Image = None
    ExifTags = None

try:
    import piexif
except ImportError:
    piexif = None

try:
    import mutagen
    from mutagen.mp3 import MP3
    from mutagen.id3 import ID3
    from mutagen.flac import FLAC
    from mutagen.mp4 import MP4
except ImportError:
    mutagen = None
    MP3 = None
    ID3 = None
    FLAC = None
    MP4 = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import yaml
except ImportError:
    yaml = None

try:
    from hachoir.metadata import extractMetadata
    from hachoir.parser import createParser
except ImportError:
    extractMetadata = None
    createParser = None

from src.config.logging_config import get_logger
from src.core.database import db

logger = get_logger("core.extractor")


def _format_size(size_bytes: int | float) -> str:
    """Format byte count into human-readable size string."""
    try:
        val = float(size_bytes)
    except (ValueError, TypeError):
        return "0 B"
    if val < 0:
        return "0 B"
    for unit in ["B", "KB", "MB", "GB", "TB"]:
        if val < 1024.0 or unit == "TB":
            if unit == "B":
                return f"{int(val)} B"
            return f"{val:.1f} {unit}"
        val /= 1024.0
    return f"{val:.1f} TB"


def _format_duration(seconds: float | int) -> str:
    """Format duration in seconds to MM:SS or HH:MM:SS string."""
    try:
        sec = max(0, int(round(float(seconds))))
    except (ValueError, TypeError):
        return "0:00"
    m, s = divmod(sec, 60)
    h, m = divmod(m, 60)
    if h > 0:
        return f"{h}:{m:02d}:{s:02d}"
    return f"{m}:{s:02d}"


def _dms_to_decimal(degrees: Any, minutes: Any, seconds: Any, ref: str = "") -> float | None:
    """Convert Degrees, Minutes, Seconds to decimal coordinate."""
    try:
        def _to_float(val: Any) -> float:
            if hasattr(val, "numerator") and hasattr(val, "denominator"):
                return float(val.numerator) / float(val.denominator) if val.denominator else 0.0
            if isinstance(val, tuple) and len(val) == 2:
                return float(val[0]) / float(val[1]) if val[1] else 0.0
            return float(val)

        d = _to_float(degrees)
        m = _to_float(minutes)
        s = _to_float(seconds)
        dec = d + (m / 60.0) + (s / 3600.0)
        if ref.upper() in ("S", "W"):
            dec = -dec
        return round(dec, 6)
    except Exception:
        return None


def _parse_yaml_or_kv(text: str) -> dict[str, Any]:
    """Parse YAML text or fallback to simple key:value parsing."""
    if yaml is not None:
        try:
            parsed = yaml.safe_load(text)
            if isinstance(parsed, dict):
                return parsed
        except Exception:
            pass

    # Simple line-by-line fallback
    result: dict[str, Any] = {}
    for line in text.strip().splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if ":" in line:
            k, v = line.split(":", 1)
            result[k.strip()] = v.strip().strip("'\"")
    return result


class MetadataExtractor:
    """Object-oriented metadata extractor with optional DB persistence.
    
    Supports comprehensive extraction from:
    - PDF documents (PyPDF2)
    - Images (PIL, EXIF, GPS coordinates, PNG chunks, SVG)
    - Audio (Mutagen, ID3, FLAC, MP4/M4A, WAV)
    - Office documents (Word DOCX, Excel XLSX, PowerPoint PPTX, OpenDocument ODT/ODS/ODP)
    - Archives (ZIP, TAR, GZ, BZ2, XZ)
    - Structured data (JSON, CSV, TSV, XML, HTML, YAML, INI, SQL)
    - Source code & Markdown (SLOC metrics, frontmatter)
    - Databases (SQLite DB schema, tables, views, pages)
    - Video files (Hachoir / Mutagen)
    - Text files & generic binary fallback
    """

    IMAGE_EXTENSIONS = {
        ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".tif", ".webp", ".ico", ".svg"
    }
    AUDIO_EXTENSIONS = {
        ".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg", ".opus", ".wma", ".aiff", ".aif"
    }
    OFFICE_EXTENSIONS = {
        ".docx", ".doc", ".xlsx", ".xlsm", ".xltx", ".xls", ".pptx", ".ppt", ".odt", ".ods", ".odp"
    }
    ARCHIVE_EXTENSIONS = {
        ".zip", ".tar", ".gz", ".tgz", ".bz2", ".tbz2", ".xz", ".txz", ".7z"
    }
    STRUCTURED_DATA_EXTENSIONS = {
        ".json", ".csv", ".tsv", ".xml", ".html", ".htm", ".xhtml", ".yaml", ".yml",
        ".ini", ".cfg", ".conf", ".toml", ".env", ".sql"
    }
    CODE_EXTENSIONS = {
        ".py", ".pyw", ".js", ".mjs", ".cjs", ".jsx", ".ts", ".tsx",
        ".c", ".cpp", ".cc", ".cxx", ".h", ".hpp", ".hxx", ".cs",
        ".java", ".go", ".rs", ".php", ".rb", ".swift", ".kt", ".kts",
        ".sh", ".bash", ".zsh", ".bat", ".cmd", ".ps1", ".lua", ".r",
        ".scala", ".dart", ".md", ".markdown", ".css", ".scss", ".less"
    }
    DATABASE_EXTENSIONS = {
        ".db", ".sqlite", ".sqlite3", ".db3", ".s3db", ".sl3"
    }
    VIDEO_EXTENSIONS = {
        ".mp4", ".mkv", ".avi", ".mov", ".wmv", ".flv", ".webm", ".m4v", ".3gp", ".mpg", ".mpeg"
    }

    def __init__(self, db_client: db.MetadataDatabase | None = None) -> None:
        self.db_client = db_client or db.db_manager

    @staticmethod
    def _validate_file_path(file_path: str) -> tuple[bool, str]:
        """Validate that the provided path exists and is a file."""
        if not file_path:
            return False, "No file path provided."
        if not os.path.exists(file_path):
            return False, f"File not found: {file_path}"
        if not os.path.isfile(file_path):
            return False, f"Path is not a file: {file_path}"
        return True, ""

    # ------------------------------------------------------------------
    # 1. PDF Metadata Extraction
    # ------------------------------------------------------------------
    def extract_pdf_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from a PDF file."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        if PdfReader is None:
            return {"Error": "PyPDF2 is not installed."}

        try:
            reader = PdfReader(file_path)
            info = reader.metadata or {}
            meta_dict: dict[str, Any] = {}
            for k, v in info.items():
                clean_key = k[1:] if k.startswith("/") else k
                if v is not None and str(v).strip():
                    meta_dict[clean_key] = str(v)

            meta_dict["Pages"] = len(reader.pages)
            meta_dict["Is Encrypted"] = "Yes" if reader.is_encrypted else "No"
            try:
                meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
            except Exception:
                pass

            logger.info("PDF metadata extracted for %s (%d pages)", file_path, len(reader.pages))
            return meta_dict
        except Exception as e:
            logger.warning("PDF extraction failed for %s: %s", file_path, e)
            return {"Error": f"PDF extraction failed: {e}"}

    # ------------------------------------------------------------------
    # 2. Image Metadata Extraction (EXIF, GPS, PNG chunks, SVG)
    # ------------------------------------------------------------------
    def extract_image_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract rich metadata from image files (JPEG, PNG, GIF, BMP, TIFF, WEBP, SVG, ICO)."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()

        # Handle SVG vector images separately
        if ext == ".svg":
            return self._extract_svg_metadata(file_path)

        if Image is None:
            return self._extract_hachoir_fallback(file_path, "Image")

        try:
            meta_dict: dict[str, Any] = {}
            with Image.open(file_path) as img:
                meta_dict["Format"] = img.format or ext.upper().lstrip(".")
                meta_dict["Mode"] = img.mode
                meta_dict["Width"] = img.width
                meta_dict["Height"] = img.height
                meta_dict["Dimensions"] = f"{img.width} x {img.height}"
                
                # Megapixels & Aspect ratio
                mp = round((img.width * img.height) / 1_000_000, 2)
                meta_dict["Megapixels"] = f"{mp} MP"
                if img.height > 0:
                    meta_dict["Aspect Ratio"] = f"{round(img.width / img.height, 2)}:1"

                # Animation frames for GIF / WEBP / APNG
                if getattr(img, "is_animated", False):
                    meta_dict["Animated"] = "Yes"
                    meta_dict["Frame Count"] = getattr(img, "n_frames", 1)

                # DPI
                dpi = img.info.get("dpi")
                if dpi:
                    if isinstance(dpi, tuple) and len(dpi) >= 2:
                        meta_dict["DPI"] = f"{int(dpi[0])} x {int(dpi[1])}"
                    else:
                        meta_dict["DPI"] = str(dpi)

                # PNG text chunks / info dictionary tags
                for key in ["Title", "Author", "Description", "Copyright", "Creation Time", "Software", "Disclaimer", "Warning", "Source", "Comment"]:
                    if key in img.info:
                        meta_dict[key] = str(img.info[key])
                    elif key.lower() in img.info:
                        meta_dict[key] = str(img.info[key.lower()])

                # EXIF Data Extraction
                exif_data = None
                try:
                    exif_data = img.getexif()
                except Exception:
                    pass

                if exif_data:
                    gps_ifd = {}
                    tag_map = ExifTags.TAGS if ExifTags else {}

                    for tag_id, value in exif_data.items():
                        tag_name = tag_map.get(tag_id, str(tag_id))
                        if tag_name == "GPSInfo" or tag_id == 0x8825:
                            if isinstance(value, dict):
                                gps_ifd = value
                            elif hasattr(exif_data, "get_ifd"):
                                try:
                                    gps_ifd = exif_data.get_ifd(0x8825)
                                except Exception:
                                    pass
                            continue

                        # Extract meaningful standard EXIF tags
                        if tag_name in [
                            "Make", "Model", "Software", "Artist", "Copyright", "ImageDescription",
                            "DateTime", "DateTimeOriginal", "DateTimeDigitized", "Orientation",
                            "ExposureTime", "FNumber", "ISOSpeedRatings", "FocalLength",
                            "Flash", "WhiteBalance", "MeteringMode", "LensModel", "LensMake",
                            "BodySerialNumber", "CameraSerialNumber"
                        ] and value:
                            meta_dict[tag_name] = str(value)

                    # Extract IFD sub-dictionaries if available
                    if hasattr(exif_data, "get_ifd"):
                        try:
                            exif_ifd = exif_data.get_ifd(ExifTags.IFD.Exif) if hasattr(ExifTags, "IFD") else {}
                            for sub_tag_id, sub_val in exif_ifd.items():
                                sub_name = tag_map.get(sub_tag_id, str(sub_tag_id))
                                if sub_name in [
                                    "DateTimeOriginal", "DateTimeDigitized", "ExposureTime",
                                    "FNumber", "ISOSpeedRatings", "FocalLength", "Flash",
                                    "WhiteBalance", "MeteringMode", "LensModel", "LensMake",
                                    "UserComment"
                                ] and sub_val and sub_name not in meta_dict:
                                    if sub_name == "UserComment" and isinstance(sub_val, bytes):
                                        try:
                                            meta_dict[sub_name] = sub_val.decode("utf-8", errors="ignore").strip("\x00")
                                        except Exception:
                                            pass
                                    else:
                                        meta_dict[sub_name] = str(sub_val)
                        except Exception:
                            pass

                        if not gps_ifd and hasattr(ExifTags, "IFD"):
                            try:
                                gps_ifd = exif_data.get_ifd(ExifTags.IFD.GPSInfo)
                            except Exception:
                                pass

                    # Parse GPS Coordinates
                    if gps_ifd:
                        gps_tag_map = ExifTags.GPSTAGS if ExifTags and hasattr(ExifTags, "GPSTAGS") else {}
                        gps_named = {gps_tag_map.get(k, str(k)): v for k, v in gps_ifd.items()}

                        lat_raw = gps_named.get("GPSLatitude")
                        lat_ref = str(gps_named.get("GPSLatitudeRef", "N"))
                        lon_raw = gps_named.get("GPSLongitude")
                        lon_ref = str(gps_named.get("GPSLongitudeRef", "E"))
                        alt_raw = gps_named.get("GPSAltitude")

                        if lat_raw and lon_raw:
                            try:
                                if isinstance(lat_raw, (tuple, list)) and len(lat_raw) == 3:
                                    dec_lat = _dms_to_decimal(lat_raw[0], lat_raw[1], lat_raw[2], lat_ref)
                                else:
                                    dec_lat = float(lat_raw)

                                if isinstance(lon_raw, (tuple, list)) and len(lon_raw) == 3:
                                    dec_lon = _dms_to_decimal(lon_raw[0], lon_raw[1], lon_raw[2], lon_ref)
                                else:
                                    dec_lon = float(lon_raw)

                                if dec_lat is not None and dec_lon is not None:
                                    meta_dict["GPS Latitude"] = f"{dec_lat}° {lat_ref}"
                                    meta_dict["GPS Longitude"] = f"{dec_lon}° {lon_ref}"
                                    meta_dict["GPS Coordinates"] = f"{dec_lat}, {dec_lon}"
                            except Exception:
                                pass

                        if alt_raw:
                            try:
                                alt_val = float(alt_raw.numerator) / float(alt_raw.denominator) if hasattr(alt_raw, "numerator") else float(alt_raw)
                                meta_dict["GPS Altitude"] = f"{round(alt_val, 2)} m"
                            except Exception:
                                pass

                        date_stamp = gps_named.get("GPSDateStamp")
                        if date_stamp:
                            meta_dict["GPS Date"] = str(date_stamp)

            try:
                meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
            except Exception:
                pass

            logger.info("Image metadata extracted successfully for %s (%d fields)", file_path, len(meta_dict))
            return meta_dict
        except Exception as e:
            logger.warning("PIL image extraction error on %s: %s", file_path, e)
            return self._extract_hachoir_fallback(file_path, "Image")

    def _extract_svg_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from SVG vector graphics file."""
        try:
            tree = ET.parse(file_path)
            root = tree.getroot()

            meta_dict: dict[str, Any] = {
                "Format": "SVG Vector Graphic",
                "Root Element": root.tag.split("}")[-1] if "}" in root.tag else root.tag,
            }

            for attr in ["width", "height", "viewBox", "version"]:
                val = root.get(attr)
                if val:
                    meta_dict[attr.capitalize() if attr != "viewBox" else "ViewBox"] = val

            # Look for <title> and <desc> tags
            for child in root:
                tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if tag.lower() == "title" and child.text:
                    meta_dict["Title"] = child.text.strip()
                elif tag.lower() == "desc" and child.text:
                    meta_dict["Description"] = child.text.strip()

            meta_dict["Total Elements"] = len(list(root.iter()))
            meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
            return meta_dict
        except Exception:
            return self.extract_text_metadata(file_path)

    # ------------------------------------------------------------------
    # 3. Audio Metadata Extraction (Mutagen, ID3, FLAC, MP4, WAV)
    # ------------------------------------------------------------------
    def extract_audio_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from audio files (MP3, WAV, FLAC, M4A, AAC, OGG, OPUS, WMA)."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()
        meta_dict: dict[str, Any] = {}

        # Wave standard library parser fallback if mutagen unavailable
        if ext == ".wav" and mutagen is None:
            return self._extract_wav_stdlib(file_path)

        if mutagen is not None:
            try:
                audio = mutagen.File(file_path)
                if audio is not None:
                    # Technical audio properties
                    if hasattr(audio, "info") and audio.info is not None:
                        info = audio.info
                        if hasattr(info, "length") and info.length:
                            meta_dict["Duration"] = _format_duration(info.length)
                            meta_dict["Duration (seconds)"] = round(info.length, 2)
                        if hasattr(info, "bitrate") and info.bitrate:
                            br_kbps = round(info.bitrate / 1000) if info.bitrate > 1000 else info.bitrate
                            meta_dict["Bitrate"] = f"{br_kbps} kbps"
                        if hasattr(info, "sample_rate") and info.sample_rate:
                            meta_dict["Sample Rate"] = f"{info.sample_rate} Hz"
                        if hasattr(info, "channels") and info.channels:
                            ch_map = {1: "1 (Mono)", 2: "2 (Stereo)", 6: "6 (5.1 Surround)"}
                            meta_dict["Channels"] = ch_map.get(info.channels, str(info.channels))
                        if hasattr(info, "bits_per_sample") and info.bits_per_sample:
                            meta_dict["Bits Per Sample"] = f"{info.bits_per_sample}-bit"

                    # Format-specific metadata tags
                    if ext == ".mp3":
                        meta_dict["Audio Format"] = "MP3 (MPEG Audio Layer III)"
                        if hasattr(audio, "tags") and audio.tags:
                            tag_id_map = {
                                "TIT2": "Title",
                                "TPE1": "Artist",
                                "TALB": "Album",
                                "TPE2": "Album Artist",
                                "TCON": "Genre",
                                "TDRC": "Date",
                                "TYER": "Year",
                                "TRCK": "Track Number",
                                "TPOS": "Disc Number",
                                "TCOM": "Composer",
                                "TSSE": "Encoder",
                                "COMM": "Comment",
                            }
                            for tag_key, tag_label in tag_id_map.items():
                                for t_k in audio.tags.keys():
                                    if t_k.startswith(tag_key):
                                        val = str(audio.tags[t_k])
                                        if val.strip():
                                            meta_dict[tag_label] = val
                                            break

                    elif ext == ".flac":
                        meta_dict["Audio Format"] = "FLAC (Free Lossless Audio Codec)"
                        if hasattr(audio, "tags") and audio.tags:
                            for k, v in audio.tags.items():
                                clean_key = k.title()
                                if isinstance(v, (list, tuple)) and v:
                                    meta_dict[clean_key] = str(v[0])
                                elif v:
                                    meta_dict[clean_key] = str(v)

                    elif ext in [".m4a", ".aac", ".mp4"]:
                        meta_dict["Audio Format"] = "M4A / AAC Audio"
                        if hasattr(audio, "tags") and audio.tags:
                            mp4_tags = {
                                "\xa9nam": "Title",
                                "\xa9ART": "Artist",
                                "\xa9alb": "Album",
                                "aART": "Album Artist",
                                "\xa9day": "Date",
                                "\xa9gen": "Genre",
                                "\xa9wrt": "Composer",
                                "\xa9cmt": "Comment",
                                "\xa9too": "Encoder",
                                "trkn": "Track Number",
                                "disk": "Disc Number",
                            }
                            for mp4_key, label in mp4_tags.items():
                                if mp4_key in audio.tags:
                                    val = audio.tags[mp4_key]
                                    if isinstance(val, (list, tuple)) and val:
                                        meta_dict[label] = str(val[0])
                                    else:
                                        meta_dict[label] = str(val)

                    elif ext in [".ogg", ".opus"]:
                        meta_dict["Audio Format"] = "OGG / Vorbis / Opus Audio"
                        if hasattr(audio, "tags") and audio.tags:
                            for k, v in audio.tags.items():
                                clean_key = k.title()
                                val_str = str(v[0]) if isinstance(v, (list, tuple)) and v else str(v)
                                if val_str.strip():
                                    meta_dict[clean_key] = val_str

                    elif ext == ".wav":
                        meta_dict["Audio Format"] = "WAV (RIFF Waveform Audio)"
                        if hasattr(audio, "tags") and audio.tags:
                            for k, v in audio.tags.items():
                                meta_dict[str(k).title()] = str(v)

                    if meta_dict:
                        meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                        return meta_dict
            except Exception as e:
                logger.warning("Mutagen audio extraction error on %s: %s", file_path, e)

        # Fallback to wave stdlib or Hachoir
        if ext == ".wav":
            return self._extract_wav_stdlib(file_path)
        return self._extract_hachoir_fallback(file_path, "Audio")

    def _extract_wav_stdlib(self, file_path: str) -> dict[str, Any]:
        """Extract WAV technical metadata using Python's standard wave module."""
        try:
            with wave.open(file_path, "rb") as wf:
                channels = wf.getnchannels()
                sample_width = wf.getsampwidth()
                framerate = wf.getframerate()
                n_frames = wf.getnframes()
                duration = n_frames / float(framerate) if framerate else 0.0

                ch_map = {1: "1 (Mono)", 2: "2 (Stereo)"}
                return {
                    "Audio Format": "WAV (RIFF Waveform Audio)",
                    "Duration": _format_duration(duration),
                    "Duration (seconds)": round(duration, 2),
                    "Sample Rate": f"{framerate} Hz",
                    "Channels": ch_map.get(channels, str(channels)),
                    "Bits Per Sample": f"{sample_width * 8}-bit",
                    "Total Frames": n_frames,
                    "File Size": _format_size(os.path.getsize(file_path)),
                }
        except Exception as e:
            return {"Error": f"WAV extraction failed: {e}"}

    # ------------------------------------------------------------------
    # 4. Office & OpenDocument Metadata (DOCX, XLSX, PPTX, ODF)
    # ------------------------------------------------------------------
    def extract_office_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from Microsoft Office (.docx, .xlsx, .pptx) and OpenDocument (.odt, .ods, .odp)."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()

        # 1. Word Document (.docx)
        if ext == ".docx":
            if docx is not None:
                try:
                    doc = docx.Document(file_path)
                    props = doc.core_properties
                    meta_dict: dict[str, Any] = {"Document Format": "Microsoft Word (DOCX)"}

                    prop_map = {
                        "Title": props.title,
                        "Subject": props.subject,
                        "Author": props.author,
                        "Keywords": props.keywords,
                        "Comments": props.comments,
                        "Last Modified By": props.last_modified_by,
                        "Category": props.category,
                        "Content Status": props.content_status,
                        "Revision": props.revision,
                    }
                    for k, v in prop_map.items():
                        if v is not None and str(v).strip():
                            meta_dict[k] = str(v)

                    if props.created:
                        meta_dict["Created"] = props.created.strftime("%Y-%m-%d %H:%M:%S")
                    if props.modified:
                        meta_dict["Modified"] = props.modified.strftime("%Y-%m-%d %H:%M:%S")
                    if props.last_printed:
                        meta_dict["Last Printed"] = props.last_printed.strftime("%Y-%m-%d %H:%M:%S")

                    meta_dict["Paragraph Count"] = len(doc.paragraphs)
                    meta_dict["Table Count"] = len(doc.tables)
                    meta_dict["Section Count"] = len(doc.sections)
                    meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                    return meta_dict
                except Exception as e:
                    logger.warning("python-docx extraction failed on %s: %s", file_path, e)

            # Fallback to OpenXML zip parsing
            return self._extract_openxml_zip(file_path, "Microsoft Word (DOCX)")

        # 2. Excel Spreadsheet (.xlsx, .xlsm, .xltx)
        if ext in [".xlsx", ".xlsm", ".xltx"]:
            if openpyxl is not None:
                try:
                    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                    props = wb.properties
                    meta_dict = {"Document Format": "Microsoft Excel (XLSX)"}

                    prop_map = {
                        "Title": props.title,
                        "Subject": props.subject,
                        "Creator": props.creator,
                        "Keywords": props.keywords,
                        "Description": props.description,
                        "Last Modified By": props.lastModifiedBy,
                        "Category": props.category,
                        "Content Status": props.contentStatus,
                        "Revision": props.version,
                    }
                    for k, v in prop_map.items():
                        if v is not None and str(v).strip():
                            meta_dict[k] = str(v)

                    if props.created:
                        meta_dict["Created"] = props.created.strftime("%Y-%m-%d %H:%M:%S")
                    if props.modified:
                        meta_dict["Modified"] = props.modified.strftime("%Y-%m-%d %H:%M:%S")

                    meta_dict["Total Sheets"] = len(wb.sheetnames)
                    meta_dict["Sheet Names"] = ", ".join(wb.sheetnames)
                    meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                    wb.close()
                    return meta_dict
                except Exception as e:
                    logger.warning("openpyxl extraction failed on %s: %s", file_path, e)

            return self._extract_openxml_zip(file_path, "Microsoft Excel (XLSX)")

        # 3. PowerPoint Presentation (.pptx)
        if ext in [".pptx", ".pptm"]:
            return self._extract_openxml_zip(file_path, "Microsoft PowerPoint (PPTX)")

        # 4. OpenDocument formats (.odt, .ods, .odp)
        if ext in [".odt", ".ods", ".odp"]:
            return self._extract_opendocument_zip(file_path)

        # Fallback to Hachoir parser for legacy binary office (.doc, .xls, .ppt)
        return self._extract_hachoir_fallback(file_path, "Office Document")

    def _extract_openxml_zip(self, file_path: str, format_name: str) -> dict[str, Any]:
        """Extract metadata directly from OpenXML package (ZIP containing core.xml and app.xml)."""
        meta_dict: dict[str, Any] = {"Document Format": format_name}
        try:
            with zipfile.ZipFile(file_path, "r") as zf:
                # 1. Parse docProps/core.xml
                if "docProps/core.xml" in zf.namelist():
                    core_xml = zf.read("docProps/core.xml")
                    root = ET.fromstring(core_xml)
                    namespaces = {
                        "dc": "http://purl.org/dc/elements/1.1/",
                        "dcterms": "http://purl.org/dc/terms/",
                        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                    }
                    elements = {
                        "Title": root.find("dc:title", namespaces),
                        "Subject": root.find("dc:subject", namespaces),
                        "Creator": root.find("dc:creator", namespaces),
                        "Keywords": root.find("cp:keywords", namespaces),
                        "Description": root.find("dc:description", namespaces),
                        "Last Modified By": root.find("cp:lastModifiedBy", namespaces),
                        "Revision": root.find("cp:revision", namespaces),
                        "Category": root.find("cp:category", namespaces),
                        "Created": root.find("dcterms:created", namespaces),
                        "Modified": root.find("dcterms:modified", namespaces),
                    }
                    for label, elem in elements.items():
                        if elem is not None and elem.text and elem.text.strip():
                            meta_dict[label] = elem.text.strip()

                # 2. Parse docProps/app.xml
                if "docProps/app.xml" in zf.namelist():
                    app_xml = zf.read("docProps/app.xml")
                    root = ET.fromstring(app_xml)
                    for child in root:
                        tag_name = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                        if tag_name in [
                            "Application", "AppVersion", "Company", "TotalTime", "Pages",
                            "Words", "Characters", "Slides", "Paragraphs", "Lines",
                            "PresentationFormat", "Notes", "HiddenSlides"
                        ] and child.text and child.text.strip():
                            meta_dict[tag_name] = child.text.strip()

            meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
            return meta_dict
        except Exception:
            return self._extract_hachoir_fallback(file_path, format_name)

    def _extract_opendocument_zip(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from OpenDocument format package (.odt, .ods, .odp)."""
        ext = os.path.splitext(file_path)[1].lower()
        fmt_map = {".odt": "OpenDocument Text (ODT)", ".ods": "OpenDocument Spreadsheet (ODS)", ".odp": "OpenDocument Presentation (ODP)"}
        meta_dict: dict[str, Any] = {"Document Format": fmt_map.get(ext, "OpenDocument")}

        try:
            with zipfile.ZipFile(file_path, "r") as zf:
                if "meta.xml" in zf.namelist():
                    meta_xml = zf.read("meta.xml")
                    root = ET.fromstring(meta_xml)
                    for elem in root.iter():
                        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
                        if tag in ["title", "description", "subject", "creator", "initial-creator", "creation-date", "date", "editing-cycles", "editing-duration"] and elem.text and elem.text.strip():
                            label = tag.replace("-", " ").title()
                            meta_dict[label] = elem.text.strip()
                        elif tag == "document-statistic":
                            for k, v in elem.attrib.items():
                                clean_k = k.split("}")[-1] if "}" in k else k
                                meta_dict[clean_k.replace("-", " ").title()] = str(v)

            meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
            return meta_dict
        except Exception:
            return self._extract_hachoir_fallback(file_path, "OpenDocument")

    # ------------------------------------------------------------------
    # 5. Archive Metadata Extraction (ZIP, TAR, GZ, BZ2, XZ)
    # ------------------------------------------------------------------
    def extract_archive_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from archive files (ZIP, TAR, GZ, TGZ, BZ2, XZ)."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()
        full_name = os.path.basename(file_path).lower()

        # Handle ZIP archives
        if ext == ".zip":
            try:
                with zipfile.ZipFile(file_path, "r") as zf:
                    infos = zf.infolist()
                    total_uncompressed = sum(info.file_size for info in infos)
                    total_compressed = sum(info.compress_size for info in infos)
                    file_count = sum(1 for info in infos if not info.is_dir())
                    dir_count = sum(1 for info in infos if info.is_dir())
                    is_encrypted = any(info.flag_bits & 0x1 for info in infos)

                    ratio = (1.0 - (total_compressed / total_uncompressed)) * 100.0 if total_uncompressed else 0.0
                    comment = zf.comment.decode("utf-8", errors="ignore") if zf.comment else ""

                    meta_dict = {
                        "Archive Format": "ZIP Archive",
                        "Total Files": file_count,
                        "Total Directories": dir_count,
                        "Uncompressed Size": _format_size(total_uncompressed),
                        "Compressed Size": _format_size(total_compressed),
                        "Compression Ratio": f"{max(0.0, round(ratio, 1))}%",
                        "Encrypted / Password Protected": "Yes" if is_encrypted else "No",
                    }
                    if comment:
                        meta_dict["Archive Comment"] = comment

                    # Sample contained files
                    sample_files = [info.filename for info in infos[:10] if not info.is_dir()]
                    if sample_files:
                        meta_dict["Sample Files"] = ", ".join(sample_files)

                    meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                    return meta_dict
            except Exception as e:
                return {"Error": f"ZIP archive extraction failed: {e}"}

        # Handle TAR / TAR.GZ / TAR.BZ2 / TAR.XZ
        if ext in [".tar", ".tgz", ".tbz2", ".txz"] or full_name.endswith((".tar.gz", ".tar.bz2", ".tar.xz")):
            try:
                with tarfile.open(file_path, "r:*") as tf:
                    members = tf.getmembers()
                    file_count = sum(1 for m in members if m.isfile())
                    dir_count = sum(1 for m in members if m.isdir())
                    total_size = sum(m.size for m in members if m.isfile())

                    meta_dict = {
                        "Archive Format": "TAR Archive Container",
                        "Total Files": file_count,
                        "Total Directories": dir_count,
                        "Uncompressed Size": _format_size(total_size),
                    }
                    sample_files = [m.name for m in members[:10] if m.isfile()]
                    if sample_files:
                        meta_dict["Sample Files"] = ", ".join(sample_files)

                    meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                    return meta_dict
            except Exception as e:
                return {"Error": f"TAR archive extraction failed: {e}"}

        # Generic archive fallback
        return self.extract_generic_metadata(file_path)

    # ------------------------------------------------------------------
    # 6. Structured Data & Config Metadata (JSON, CSV, XML, YAML, INI, SQL)
    # ------------------------------------------------------------------
    def extract_structured_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from structured data files (JSON, CSV, TSV, XML, HTML, YAML, INI, SQL)."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()

        # 1. JSON Data
        if ext == ".json":
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read(1024 * 1024)
                    data = json.loads(content)

                stat = os.stat(file_path)
                meta_dict = {
                    "Data Format": "JSON (JavaScript Object Notation)",
                    "Root Type": "Object (Dictionary)" if isinstance(data, dict) else ("Array (List)" if isinstance(data, list) else type(data).__name__),
                    "File Size": _format_size(stat.st_size),
                }

                if isinstance(data, dict):
                    meta_dict["Top-Level Keys Count"] = len(data)
                    keys_sample = list(data.keys())[:12]
                    meta_dict["Top-Level Keys"] = ", ".join(str(k) for k in keys_sample)
                elif isinstance(data, list):
                    meta_dict["Total Items"] = len(data)

                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    meta_dict["Line Count"] = sum(1 for _ in f)

                return meta_dict
            except Exception:
                return self.extract_text_metadata(file_path)

        # 2. CSV / TSV Data
        if ext in [".csv", ".tsv"]:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    sample = f.read(8192)
                    f.seek(0)
                    dialect = None
                    try:
                        dialect = csv.Sniffer().sniff(sample)
                    except Exception:
                        pass

                    delimiter = dialect.delimiter if dialect else ("\t" if ext == ".tsv" else ",")
                    reader = csv.reader(f, delimiter=delimiter)
                    first_row = next(reader, None)
                    row_count = 1 + sum(1 for _ in reader) if first_row else 0

                stat = os.stat(file_path)
                delim_name = {",": "Comma (,)", "\t": "Tab (\\t)", ";": "Semicolon (;)", "|": "Pipe (|)"}.get(delimiter, delimiter)

                meta_dict = {
                    "Data Format": "TSV (Tab-Separated Values)" if ext == ".tsv" else "CSV (Comma-Separated Values)",
                    "Delimiter": delim_name,
                    "Total Rows": row_count,
                    "Total Columns": len(first_row) if first_row else 0,
                    "File Size": _format_size(stat.st_size),
                }
                if first_row:
                    header_preview = [str(col).strip() for col in first_row[:10] if str(col).strip()]
                    if header_preview:
                        meta_dict["Header Columns"] = ", ".join(header_preview)

                return meta_dict
            except Exception:
                return self.extract_text_metadata(file_path)

        # 3. XML / HTML Data
        if ext in [".xml", ".html", ".htm", ".xhtml"]:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text_content = f.read()

                meta_dict = {
                    "Data Format": "HTML Document" if ext in [".html", ".htm"] else "XML Document",
                    "File Size": _format_size(os.path.getsize(file_path)),
                    "Line Count": text_content.count("\n") + 1,
                }

                # Try finding <title> tag
                title_match = re.search(r"<title[^>]*>(.*?)</title>", text_content, re.IGNORECASE | re.DOTALL)
                if title_match:
                    meta_dict["Title"] = title_match.group(1).strip()

                try:
                    tree = ET.fromstring(text_content)
                    root_tag = tree.tag.split("}")[-1] if "}" in tree.tag else tree.tag
                    meta_dict["Root Element"] = root_tag
                    meta_dict["Total Elements"] = len(list(tree.iter()))
                except Exception:
                    pass

                return meta_dict
            except Exception:
                return self.extract_text_metadata(file_path)

        # 4. YAML Data
        if ext in [".yaml", ".yml"]:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
                data = _parse_yaml_or_kv(text)
                meta_dict = {
                    "Data Format": "YAML Document",
                    "File Size": _format_size(os.path.getsize(file_path)),
                    "Line Count": text.count("\n") + 1,
                }
                if data and isinstance(data, dict):
                    meta_dict["Top-Level Keys"] = ", ".join(str(k) for k in list(data.keys())[:10])
                return meta_dict
            except Exception:
                return self.extract_text_metadata(file_path)

        # 5. INI / Configuration Data
        if ext in [".ini", ".cfg", ".conf", ".env", ".properties", ".toml"]:
            try:
                cp = configparser.ConfigParser()
                cp.read(file_path, encoding="utf-8")
                sections = cp.sections()
                meta_dict = {
                    "Data Format": "Configuration File",
                    "Section Count": len(sections),
                    "Sections": ", ".join(sections[:10]) if sections else "None",
                    "File Size": _format_size(os.path.getsize(file_path)),
                }
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    meta_dict["Line Count"] = sum(1 for _ in f)
                return meta_dict
            except Exception:
                return self.extract_text_metadata(file_path)

        # 6. SQL Script Data
        if ext == ".sql":
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    lines = f.readlines()
                text = "".join(lines)
                tables = re.findall(r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?([`\"'\[]?\w+[`\"'\]]?)", text, re.IGNORECASE)
                clean_tables = [t.strip("`\"'[]") for t in tables]

                meta_dict = {
                    "Data Format": "SQL Database Script",
                    "Line Count": len(lines),
                    "Tables Defined": ", ".join(clean_tables[:10]) if clean_tables else "None",
                    "Table Count": len(clean_tables),
                    "File Size": _format_size(os.path.getsize(file_path)),
                }
                return meta_dict
            except Exception:
                return self.extract_text_metadata(file_path)

        return self.extract_text_metadata(file_path)

    # ------------------------------------------------------------------
    # 7. Source Code & Markdown Metadata
    # ------------------------------------------------------------------
    def extract_code_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract source lines of code (SLOC) and language metadata from code/markdown."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()

        # Language mapping
        lang_map = {
            ".py": "Python", ".pyw": "Python", ".js": "JavaScript", ".mjs": "JavaScript",
            ".cjs": "JavaScript", ".jsx": "React JSX", ".ts": "TypeScript", ".tsx": "React TSX",
            ".c": "C", ".cpp": "C++", ".cc": "C++", ".cxx": "C++", ".h": "C/C++ Header",
            ".hpp": "C++ Header", ".cs": "C#", ".java": "Java", ".go": "Go", ".rs": "Rust",
            ".php": "PHP", ".rb": "Ruby", ".swift": "Swift", ".kt": "Kotlin", ".sh": "Shell Script",
            ".bash": "Bash Script", ".bat": "Batch Script", ".ps1": "PowerShell", ".lua": "Lua",
            ".r": "R", ".scala": "Scala", ".dart": "Dart", ".css": "CSS", ".scss": "SCSS",
            ".md": "Markdown", ".markdown": "Markdown"
        }

        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                lines = f.readlines()

            total_lines = len(lines)
            blank_lines = sum(1 for line in lines if not line.strip())
            comment_lines = 0

            # Estimate comments based on language syntax
            if ext in [".py", ".pyw", ".sh", ".bash", ".rb", ".r", ".ps1"]:
                comment_lines = sum(1 for line in lines if line.strip().startswith("#"))
            elif ext in [".c", ".cpp", ".cc", ".cxx", ".h", ".hpp", ".cs", ".java", ".js", ".ts", ".jsx", ".tsx", ".go", ".rs", ".swift", ".kt", ".css", ".scss"]:
                comment_lines = sum(1 for line in lines if line.strip().startswith("//") or line.strip().startswith("/*") or line.strip().startswith("*"))

            code_lines = max(0, total_lines - blank_lines - comment_lines)
            stat = os.stat(file_path)

            meta_dict: dict[str, Any] = {
                "Language": lang_map.get(ext, "Source Code"),
                "Total Lines": total_lines,
                "Code Lines (SLOC)": code_lines,
                "Comment Lines": comment_lines,
                "Blank Lines": blank_lines,
                "File Size": _format_size(stat.st_size),
                "Encoding": "utf-8",
            }

            # Extract Markdown YAML frontmatter if available
            if ext in [".md", ".markdown"] and lines:
                if lines[0].strip() == "---":
                    frontmatter_lines = []
                    for line in lines[1:]:
                        if line.strip() == "---":
                            break
                        frontmatter_lines.append(line)
                    if frontmatter_lines:
                        fm_data = _parse_yaml_or_kv("".join(frontmatter_lines))
                        if isinstance(fm_data, dict):
                            for k, v in fm_data.items():
                                meta_dict[f"Frontmatter {k.title()}"] = str(v)

                # Markdown specific metrics
                text_content = "".join(lines)
                meta_dict["Header Count"] = len(re.findall(r"^#{1,6}\s+", text_content, re.MULTILINE))
                meta_dict["Word Count"] = len(re.findall(r"\b\w+\b", text_content))
                meta_dict["Code Blocks Count"] = len(re.findall(r"```", text_content)) // 2

            return meta_dict
        except Exception:
            return self.extract_text_metadata(file_path)

    # ------------------------------------------------------------------
    # 8. Database Metadata Extraction (SQLite)
    # ------------------------------------------------------------------
    def extract_database_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract schema structure and metrics from SQLite database files."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        try:
            db_uri = f"file:{os.path.abspath(file_path)}?mode=ro"
            conn = sqlite3.connect(db_uri, uri=True)
            cursor = conn.cursor()

            try:
                # PRAGMA queries
                cursor.execute("PRAGMA page_size")
                page_size = cursor.fetchone()[0]
                cursor.execute("PRAGMA page_count")
                page_count = cursor.fetchone()[0]
                cursor.execute("PRAGMA user_version")
                user_version = cursor.fetchone()[0]
                cursor.execute("PRAGMA schema_version")
                schema_version = cursor.fetchone()[0]

                # Query tables and views
                cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")
                tables = [row[0] for row in cursor.fetchall()]

                cursor.execute("SELECT name FROM sqlite_master WHERE type='view' ORDER BY name")
                views = [row[0] for row in cursor.fetchall()]

                cursor.execute("SELECT name FROM sqlite_master WHERE type='index' AND name NOT LIKE 'sqlite_%' ORDER BY name")
                indexes = [row[0] for row in cursor.fetchall()]
            finally:
                conn.close()

            stat = os.stat(file_path)
            meta_dict = {
                "Database Engine": "SQLite 3",
                "Database Size": _format_size(stat.st_size),
                "Page Size": f"{page_size} bytes",
                "Page Count": page_count,
                "Table Count": len(tables),
                "Tables": ", ".join(tables) if tables else "None",
                "View Count": len(views),
                "Views": ", ".join(views) if views else "None",
                "Index Count": len(indexes),
                "User Version": user_version,
                "Schema Version": schema_version,
            }
            return meta_dict
        except Exception:
            return self.extract_generic_metadata(file_path)

    # ------------------------------------------------------------------
    # 9. Video Metadata Extraction
    # ------------------------------------------------------------------
    def extract_video_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from video files (MP4, MKV, AVI, MOV, WMV, FLV, WEBM)."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        ext = os.path.splitext(file_path)[1].lower()

        # Try Mutagen for MP4 / MOV / M4V containers
        if ext in [".mp4", ".mov", ".m4v"] and MP4 is not None:
            try:
                video = MP4(file_path)
                meta_dict = {"Video Format": f"{ext.upper().lstrip('.')} Video Container"}
                if video.info:
                    if video.info.length:
                        meta_dict["Duration"] = _format_duration(video.info.length)
                        meta_dict["Duration (seconds)"] = round(video.info.length, 2)
                    if video.info.bitrate:
                        meta_dict["Bitrate"] = f"{round(video.info.bitrate / 1000)} kbps"

                if video.tags:
                    tag_map = {
                        "\xa9nam": "Title",
                        "\xa9ART": "Artist",
                        "\xa9alb": "Album",
                        "\xa9day": "Date",
                        "\xa9gen": "Genre",
                        "\xa9too": "Software",
                        "\xa9cmt": "Comment",
                    }
                    for k, label in tag_map.items():
                        if k in video.tags:
                            val = video.tags[k]
                            meta_dict[label] = str(val[0]) if isinstance(val, (list, tuple)) and val else str(val)

                if len(meta_dict) > 1:
                    meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                    return meta_dict
            except Exception:
                pass

        return self._extract_hachoir_fallback(file_path, "Video")

    # ------------------------------------------------------------------
    # 10. Plain Text Metadata
    # ------------------------------------------------------------------
    def extract_text_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from a text file."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                lines = f.readlines()
            stat = os.stat(file_path)
            logger.info("Text metadata extracted for %s (%d lines)", file_path, len(lines))
            return {
                "File Size": _format_size(stat.st_size),
                "File Size (bytes)": stat.st_size,
                "Line Count": len(lines),
                "Encoding": "utf-8",
            }
        except Exception as e:
            logger.warning("Text extraction failed for %s: %s", file_path, e)
            return {"Error": f"Text extraction failed: {e}"}

    # ------------------------------------------------------------------
    # 11. Generic Binary & Fallback Metadata
    # ------------------------------------------------------------------
    def extract_generic_metadata(self, file_path: str) -> dict[str, Any]:
        """Extract generic properties, checksums, and timestamps for any file."""
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            return {"Error": error}

        try:
            stat = os.stat(file_path)
            mime_type, _ = mimetypes.guess_type(file_path)

            sha256 = hashlib.sha256()
            md5 = hashlib.md5()
            with open(file_path, "rb") as f:
                while chunk := f.read(65536):
                    sha256.update(chunk)
                    md5.update(chunk)

            created_dt = datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S")
            modified_dt = datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            accessed_dt = datetime.fromtimestamp(stat.st_atime).strftime("%Y-%m-%d %H:%M:%S")

            return {
                "File Name": os.path.basename(file_path),
                "MIME Type": mime_type or "application/octet-stream",
                "File Size": _format_size(stat.st_size),
                "File Size (bytes)": stat.st_size,
                "SHA-256 Checksum": sha256.hexdigest(),
                "MD5 Checksum": md5.hexdigest(),
                "Created Date": created_dt,
                "Modified Date": modified_dt,
                "Accessed Date": accessed_dt,
            }
        except Exception as e:
            return {"Error": f"Generic extraction failed: {e}"}

    def _extract_hachoir_fallback(self, file_path: str, category_hint: str = "") -> dict[str, Any]:
        """Internal helper to attempt Hachoir parsing with generic fallback."""
        if createParser is not None and extractMetadata is not None:
            try:
                parser = createParser(file_path)
                if parser:
                    metadata = extractMetadata(parser)
                    if metadata:
                        meta_dict: dict[str, Any] = {}
                        for item in metadata.exportPlaintext():
                            if ": " in item:
                                key, value = item.split(": ", 1)
                                meta_dict[key.strip()] = value.strip()
                        if meta_dict:
                            try:
                                meta_dict["File Size"] = _format_size(os.path.getsize(file_path))
                            except Exception:
                                pass
                            return meta_dict
            except Exception:
                pass

        generic = self.extract_generic_metadata(file_path)
        if category_hint:
            generic["Category"] = category_hint
        return generic

    # ------------------------------------------------------------------
    # Master Extraction Router
    # ------------------------------------------------------------------
    def extract(self, file_path: str) -> dict[str, Any]:
        """Extract metadata from any supported file type.
        
        Automatically detects file type and routes to the specialized extractor:
        - PDF: extract_pdf_metadata
        - Images: extract_image_metadata
        - Audio: extract_audio_metadata
        - Office: extract_office_metadata
        - Archives: extract_archive_metadata
        - Structured Data: extract_structured_metadata
        - Source Code: extract_code_metadata
        - Database: extract_database_metadata
        - Video: extract_video_metadata
        - Text: extract_text_metadata
        - Fallback: extract_generic_metadata
        """
        is_valid, error = self._validate_file_path(file_path)
        if not is_valid:
            logger.debug("File validation failed for %s: %s", file_path, error)
            return {"Error": error}

        mime_type, _ = mimetypes.guess_type(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        # 1. PDF
        if ext == ".pdf" or (mime_type and mime_type == "application/pdf"):
            return self.extract_pdf_metadata(file_path)

        # 2. Images
        if ext in self.IMAGE_EXTENSIONS or (mime_type and mime_type.startswith("image")):
            return self.extract_image_metadata(file_path)

        # 3. Audio
        if ext in self.AUDIO_EXTENSIONS or (mime_type and mime_type.startswith("audio")):
            return self.extract_audio_metadata(file_path)

        # 4. Office & OpenDocument
        if ext in self.OFFICE_EXTENSIONS:
            return self.extract_office_metadata(file_path)

        # 5. Archives
        if ext in self.ARCHIVE_EXTENSIONS or file_path.lower().endswith((".tar.gz", ".tar.bz2", ".tar.xz")):
            return self.extract_archive_metadata(file_path)

        # 6. Structured Data & Config
        if ext in self.STRUCTURED_DATA_EXTENSIONS:
            return self.extract_structured_metadata(file_path)

        # 7. Source Code & Markdown
        if ext in self.CODE_EXTENSIONS:
            return self.extract_code_metadata(file_path)

        # 8. Databases
        if ext in self.DATABASE_EXTENSIONS:
            return self.extract_database_metadata(file_path)

        # 9. Videos
        if ext in self.VIDEO_EXTENSIONS or (mime_type and mime_type.startswith("video")):
            return self.extract_video_metadata(file_path)

        # 10. Text files
        if mime_type and mime_type.startswith("text"):
            return self.extract_text_metadata(file_path)

        # 11. Generic fallback with Hachoir
        return self._extract_hachoir_fallback(file_path)

    def extract_and_store(self, file_path: str) -> tuple[dict[str, Any], Any]:
        """Extract metadata from file and store in database."""
        metadata = self.extract(file_path)
        db_row = None

        if not metadata or not isinstance(metadata, dict):
            return {"Error": "Extraction returned no metadata."}, None

        if "Error" in metadata:
            return metadata, None

        try:
            db_row = self.db_client.insert_metadata(file_path, metadata)
            logger.info("Stored metadata in DB for %s (row id=%s)", file_path, db_row[0] if db_row else None)
        except Exception as exc:
            logger.error("Failed to persist metadata for %s: %s", file_path, exc)
            metadata = {**metadata, "Error": f"Failed to persist metadata: {exc}"}
            db_row = None

        return metadata, db_row

    def batch_extract(self, file_paths: Iterable[str], progress_callback: Callable[[str, float], None] | None = None) -> dict[str, Any]:
        """Extract metadata from multiple files with optional progress reporting."""
        successful_extractions = 0
        failed_extractions = 0
        results = []
        file_paths = list(file_paths)
        total_files = len(file_paths)

        safe_progress_callback = progress_callback
        for i, file_path in enumerate(file_paths):
            try:
                if safe_progress_callback:
                    progress = (i / total_files) * 100 if total_files else 0
                    try:
                        safe_progress_callback(f"Processing: {os.path.basename(file_path)}", progress)
                    except Exception:
                        safe_progress_callback = None

                meta_dict = self.extract(file_path)

                if meta_dict and "Error" not in meta_dict:
                    row = self.db_client.insert_metadata(file_path, meta_dict)
                    results.append({"file_path": file_path, "status": "success", "data": row})
                    successful_extractions += 1
                else:
                    results.append({"file_path": file_path, "status": "failed", "error": meta_dict.get("Error", "Unknown error")})
                    failed_extractions += 1

            except Exception as e:
                results.append({"file_path": file_path, "status": "failed", "error": str(e)})
                failed_extractions += 1

        if safe_progress_callback:
            try:
                safe_progress_callback("Batch extraction completed", 100)
            except Exception:
                pass

        return {"successful": successful_extractions, "failed": failed_extractions, "total": total_files, "results": results}


_extractor = MetadataExtractor()


# Module-level wrapper functions for backward compatibility & direct access
def extract_pdf_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a PDF file."""
    return _extractor.extract_pdf_metadata(file_path)


def extract_image_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from an image file."""
    return _extractor.extract_image_metadata(file_path)


def extract_audio_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from an audio file."""
    return _extractor.extract_audio_metadata(file_path)


def extract_office_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a Microsoft Office or OpenDocument file."""
    return _extractor.extract_office_metadata(file_path)


def extract_archive_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from an archive file."""
    return _extractor.extract_archive_metadata(file_path)


def extract_structured_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a structured data / configuration file."""
    return _extractor.extract_structured_metadata(file_path)


def extract_code_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a source code or markdown file."""
    return _extractor.extract_code_metadata(file_path)


def extract_database_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a database file."""
    return _extractor.extract_database_metadata(file_path)


def extract_video_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a video file."""
    return _extractor.extract_video_metadata(file_path)


def extract_text_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from a text file."""
    return _extractor.extract_text_metadata(file_path)


def extract_generic_metadata(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract generic metadata and checksums from any file."""
    return _extractor.extract_generic_metadata(file_path)


def extract(file_path: str) -> dict[str, Any]:
    """Wrapper: Extract metadata from any supported file type."""
    return _extractor.extract(file_path)


def extract_and_store(file_path: str) -> tuple[dict[str, Any], Any]:
    """Wrapper: Extract metadata and store in database."""
    return _extractor.extract_and_store(file_path)


def batch_extract(file_paths: Iterable[str], progress_callback: Callable[[str, float], None] | None = None) -> dict[str, Any]:
    """Wrapper: Extract metadata from multiple files with optional progress reporting."""
    return _extractor.batch_extract(file_paths, progress_callback)
